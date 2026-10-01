# -*- coding: utf-8 -*-
"""
generar_informe_docx.py

Construye el informe del P1.3 en formato Word, siguiendo la misma estructura
y el mismo formato del informe entregado en el P1.2. Toma las cifras de los
resultados medidos y las capturas de la aplicacion desde la carpeta
resultados.

Uso:
    python generar_informe_docx.py
"""

import csv
import json
import os

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.shared import Inches, Pt

import tildes

SALIDA = 'resultados'
DESTINO = 'P1.3_Colaborativo_VivancoAlex.docx'
FECHA = '22 de septiembre de 2026'

# ---------------------------------------------------------------------------
# Mediciones. Todas salen de corridas reales; ninguna esta inventada.
#
# Las del recorte de 500 MB se leen del CSV y del JSON que deja capturas.py,
# que son los de la misma ejecucion de la que salieron las imagenes. Se hace
# asi y no copiando las cifras a mano para que el documento no pueda acabar
# diciendo una cosa mientras sus propias capturas dicen otra.
# ---------------------------------------------------------------------------


def cargar_filas():
    """Lee las mediciones del CSV que dejo la ultima corrida capturada."""
    ruta = os.path.join(SALIDA, 'benchmark.csv')
    if not os.path.exists(ruta):
        raise SystemExit(
            'Falta %s. Ejecuta primero:\n'
            '    python capturas.py <archivo.fna> --completo' % ruta)
    with open(ruta, encoding='utf-8') as f:
        return list(csv.DictReader(f))


def cargar_datos():
    """Lee el JSON con las mediciones que no caben en el CSV de filas."""
    ruta = os.path.join(SALIDA, 'datos_informe.json')
    if not os.path.exists(ruta):
        raise SystemExit(
            'Falta %s. Ejecuta primero:\n'
            '    python capturas.py <archivo.fna> --completo' % ruta)
    with open(ruta, encoding='utf-8') as f:
        return json.load(f)


def num(fila, clave, por_defecto=0.0):
    """Valor numerico de una celda del CSV, o el valor por defecto si falta."""
    texto = (fila.get(clave) or '').strip()
    if not texto:
        return por_defecto
    try:
        return float(texto)
    except ValueError:
        return por_defecto


FILAS = cargar_filas()
DATOS = cargar_datos()


def por_modo(*modos):
    """Filas del CSV que pertenecen a alguno de los modos indicados."""
    return [f for f in FILAS if f.get('modo') in modos]

# Conteo del genoma humano completo GRCh38.p14 (3,11 GB).
GENOMA = {
    'A': 923117203, 'C': 642552917, 'G': 645231996, 'T': 925917038,
    'N': 161611379, 'iupac': 103, 'invalidos': 0,
    'mb': 3185.0,
}
GENOMA['bases'] = GENOMA['A'] + GENOMA['C'] + GENOMA['G'] + GENOMA['T']

# Modos sobre el genoma completo: (etiqueta, tiempo, MB/s, % CPU, solape).
# El genoma no se mide desde la aplicacion porque cada pasada tarda demasiado
# para una sesion interactiva, asi que estas cifras provienen de la corrida
# por linea de comandos y se transcriben aqui. Todas son posteriores a la
# correccion del error de borde descrito en la seccion de analisis.
MODOS_GENOMA = [
    ('CPU paralelo x12', 9.323, 341.6, None, None),
    ('GPU sola', 1.488, 2140.6, None, None),
    ('CPU x4 + GPU', 2.870, 1109.6, 21.0, 2.73),
    ('CPU x8 + GPU', 3.326, 957.7, 32.0, 3.01),
    ('CPU x12 + GPU', 5.177, 615.2, 52.0, 4.70),
]

# Comparativa de modos y barridos sobre el recorte de 500 MB, leidos del CSV.
MODOS_500 = por_modo('cpu_seq', 'cpu_par', 'gpu', 'hibrido')
PROCESOS_500 = por_modo('barrido_procesos')
TROZOS_500 = por_modo('barrido_trozo')
ESCALA_500 = por_modo('escala')
MB_500 = (num(MODOS_500[0], 'bytes') / (1024 * 1024)) if MODOS_500 else 500.0

# Escalabilidad reorganizada por tamano: {tamano: {plataforma: tiempo}}
ESCALABILIDAD = []
for _fila in ESCALA_500:
    _etiqueta = _fila['etiqueta']
    if not ESCALABILIDAD or ESCALABILIDAD[-1][0] != _etiqueta:
        ESCALABILIDAD.append((_etiqueta, {}))
    ESCALABILIDAD[-1][1][_fila['plataforma']] = num(_fila, 'tiempo_s')

# Mediciones de la corrida completa que no caben en el CSV de filas.
CENSO_BYTES = DATOS.get('censo_sm') or []
_TOTAL_SM = sum(CENSO_BYTES) or 1
CENSO_SM = [100.0 * v / _TOTAL_SM for v in CENSO_BYTES]

_REC = DATOS.get('recursos') or {}
_KER = DATOS.get('recursos_kernel') or {}
NUCLEOS_CPU = _REC.get('cpu_por_nucleo') or []

RECURSOS = {
    'cpu_medio': (sum(NUCLEOS_CPU) / len(NUCLEOS_CPU)) if NUCLEOS_CPU else 0.0,
    'cpu_min': min(NUCLEOS_CPU) if NUCLEOS_CPU else 0.0,
    'cpu_max': max(NUCLEOS_CPU) if NUCLEOS_CPU else 0.0,
    'cpu_temp_media': _REC.get('cpu_temp_medio', 0.0),
    'cpu_temp_max': _REC.get('cpu_temp_max', 0.0),
    'gpu_temp_media': _REC.get('gpu_temp_medio', 0.0),
    'gpu_temp_max': _REC.get('gpu_temp_max', 0.0),
    'ram_pct': _REC.get('ram_medio', 0.0),
    'proc_mb': int(_REC.get('proc_mb_max', 0)),
    'vram_mb': int(_REC.get('vram_mb_max', 0)),
    'vram_total': 6140,
    'registros': _KER.get('registros_por_hilo', 0),
    'compartida': _KER.get('compartida_por_bloque', 0),
    'warps': _KER.get('warps_activos', 0),
    'warps_max': _KER.get('warps_maximos', 0),
    'ocupacion': _KER.get('ocupacion_pct', 0.0),
}

# Las tres respuestas declarativas de la rubrica, por nivel.
NIVELES = [
    ('Programa',
     'motor_cpu.py. La funcion contar_paralelo reparte rangos de bytes entre '
     'procesos y contar_rango cuenta cada uno con numpy.bincount.',
     'motor_gpu.py. Tres kernels encadenados: marcar las cabeceras, borrarlas '
     'dentro de la tarjeta y calcular el histograma de 256 bins.',
     'motor_gpu.py, kernel k_histograma. Cada hilo acumula en la memoria '
     'compartida de su propio multiprocesador y anota en cual se ejecuto.'),
    ('Libreria',
     'multiprocessing, de la biblioteca estandar, junto a numpy para el '
     'conteo vectorizado dentro de cada proceso.',
     'numba.cuda, que compila a PTX los kernels escritos en Python y gestiona '
     'la memoria, los streams y los lanzamientos.',
     'numba.extending.intrinsic junto a llvmlite.ir, que permiten inyectar '
     'ensamblador PTX dentro del kernel.'),
    ('Instruccion',
     'multiprocessing.Pool(processes=n) y pool.starmap_async(trabajo, tareas).',
     'El decorador @cuda.jit sobre la funcion del kernel y el lanzamiento '
     'k_histograma[bloques, hilos, stream](datos, n, hist, por_sm).',
     'La instruccion PTX mov.u32 $0, %smid; leida mediante un intrinsic, y '
     'cuda.atomic.add(por_sm, smid(), vistos) para acumular su cuota.'),
]


def miles(n):
    return '{:,}'.format(n).replace(',', '.')


# Accesos cortos a las columnas que mas se citan.
def t(fila):
    return num(fila, 'tiempo_s')


def mbs(fila):
    return num(fila, 'mb_por_s')


def sp(fila):
    return num(fila, 'speedup')


def rcpu(fila):
    return num(fila, 'reparto_cpu_pct')


def sol(fila):
    return num(fila, 'solape_s')


def una(modo):
    """Primera fila de un modo concreto, o None si ese modo no se midio."""
    filas = por_modo(modo)
    return filas[0] if filas else None


def mejor(filas):
    """La configuracion mas rapida de un grupo."""
    return min(filas, key=t) if filas else None


def peor(filas):
    """La configuracion mas lenta de un grupo."""
    return max(filas, key=t) if filas else None


CPU1 = una('cpu_seq')
CPUP = una('cpu_par')
GPU = una('gpu')
HIB = una('hibrido')


class Informe:
    def __init__(self):
        self.doc = Document()
        estilo = self.doc.styles['Normal']
        estilo.font.name = 'Calibri'
        estilo.font.size = Pt(11)
        self.figura = 0
        self._encabezado_y_pie()

    # -- estructura de pagina ---------------------------------------------

    def _campo(self, parrafo, instruccion):
        """Inserta un campo de Word, como el numero de pagina.

        python-docx no expone los campos, asi que hay que construir el XML a
        mano: una marca de inicio, la instruccion y una marca de cierre.
        """
        run = parrafo.add_run()
        inicio = OxmlElement('w:fldChar')
        inicio.set(qn('w:fldCharType'), 'begin')
        texto = OxmlElement('w:instrText')
        texto.set(qn('xml:space'), 'preserve')
        texto.text = instruccion
        fin = OxmlElement('w:fldChar')
        fin.set(qn('w:fldCharType'), 'end')
        run._r.append(inicio)
        run._r.append(texto)
        run._r.append(fin)
        run.font.size = Pt(9)

    def _encabezado_y_pie(self):
        """Encabezado con el titulo y pie con autor, numeracion y fecha."""
        seccion = self.doc.sections[0]

        encabezado = seccion.header.paragraphs[0]
        encabezado.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = encabezado.add_run('P1.3 Computación CPU & GPU Paralela - '
                                 'CPU y GPU trabajando en paralelo')
        run.font.size = Pt(9)
        run.italic = True

        pie = seccion.footer.paragraphs[0]
        pie.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = pie.add_run('Alex Vivanco  |  %s  |  Pagina ' % FECHA)
        run.font.size = Pt(9)
        self._campo(pie, 'PAGE')

    # -- utilidades de escritura ------------------------------------------

    def titulo(self, texto, tamano=14, espacio_antes=14):
        # Los titulos pueden estar formulados como pregunta, asi que aqui si
        # se acentuan los interrogativos.
        texto = tildes.tildar(texto, interrogativo=True)
        p = self.doc.add_paragraph()
        p.paragraph_format.space_before = Pt(espacio_antes)
        p.paragraph_format.space_after = Pt(6)
        r = p.add_run(texto)
        r.bold = True
        r.font.size = Pt(tamano)
        return p

    def parrafo(self, texto, justificado=True):
        p = self.doc.add_paragraph(tildes.tildar(texto))
        if justificado:
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.space_after = Pt(8)
        return p

    def vineta(self, texto):
        p = self.doc.add_paragraph(tildes.tildar(texto), style='List Bullet')
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.space_after = Pt(4)
        return p

    def imagen(self, ruta, pie, ancho=6.3):
        if not os.path.exists(ruta):
            print('  AVISO: falta la imagen %s' % ruta)
            return
        self.doc.add_picture(ruta, width=Inches(ancho))
        self.doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        self.figura += 1
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(tildes.tildar('Figura %d. %s'
                                     % (self.figura, pie)))
        r.italic = True
        r.font.size = Pt(9)
        p.paragraph_format.space_after = Pt(12)

    def tabla(self, cabecera, filas, anchos=None):
        t = self.doc.add_table(rows=1, cols=len(cabecera))
        t.style = 'Light Grid Accent 1'
        for celda, texto in zip(t.rows[0].cells, cabecera):
            celda.text = ''
            r = celda.paragraphs[0].add_run(tildes.tildar(texto))
            r.bold = True
            r.font.size = Pt(9)
        for fila in filas:
            celdas = t.add_row().cells
            for celda, texto in zip(celdas, fila):
                celda.text = ''
                r = celda.paragraphs[0].add_run(tildes.tildar(str(texto)))
                r.font.size = Pt(9)
        self.doc.add_paragraph().paragraph_format.space_after = Pt(6)
        return t

    # -- contenido ---------------------------------------------------------

    def portada(self):
        for texto, tam, negrita in (
            ('PONTIFICIA UNIVERSIDAD CATOLICA DEL ECUADOR', 13, True),
            ('Facultad de Ingenieria', 12, False),
            ('Ingenieria en Sistemas y Computacion', 12, False),
        ):
            p = self.doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(tildes.tildar(texto))
            r.bold = negrita
            r.font.size = Pt(tam)

        self.doc.add_paragraph()
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run('PROYECTO P1.3')
        r.bold = True
        r.font.size = Pt(20)

        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run('Computación CPU & GPU Paralela: procesamiento '
                      'simultáneo de cadenas de ADN sobre ambas plataformas')
        r.bold = True
        r.font.size = Pt(14)

        self.doc.add_paragraph()
        for texto in ('Asignatura: Computacion Paralela, Distribuida y en la Nube',
                      'Docente: Ing. JW Condor',
                      'Autor: Alex Vivanco',
                      'Fecha de exposicion: %s' % FECHA):
            p = self.doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.add_run(tildes.tildar(texto)).font.size = Pt(11)

    def resumen(self):
        self.titulo('1. Resumen')
        self.parrafo(
            'Este informe presenta una aplicacion en Python que procesa el archivo '
            'FASTA del genoma humano GRCh38.p14, de 3,11 GB, poniendo a trabajar la '
            'CPU multinucleo y la GPU de forma simultanea sobre el mismo archivo. A '
            'diferencia del proyecto anterior, en el que ambas plataformas se median '
            'por separado para compararlas, aqui colaboran: el archivo se divide en '
            'trozos y cada plataforma toma los que puede atender mientras la otra '
            'trabaja, de modo que el reparto no se fija de antemano sino que emerge '
            'de la velocidad que cada una demuestra durante la ejecucion.')
        self.parrafo(
            'Sobre el genoma completo, la ejecucion paralela en CPU con doce procesos '
            'tomo %.3f segundos, la ejecucion en GPU tomo %.3f segundos y la ejecucion '
            'conjunta con cuatro procesos de CPU tomo %.3f segundos, con una '
            'superposicion real de %.2f segundos durante los cuales ambas plataformas '
            'estuvieron activas. Todas las configuraciones produjeron conteos '
            'identicos, verificados sobre el histograma completo de 256 valores de '
            'byte y no unicamente sobre las cuatro bases.'
            % (MODOS_GENOMA[0][1], MODOS_GENOMA[1][1], MODOS_GENOMA[2][1],
               MODOS_GENOMA[2][4]))
        self.parrafo(
            'La aplicacion instrumenta ademas el trabajo a nivel de nucleo: registra '
            'el uso de cada uno de los doce procesadores logicos de la CPU y, dentro '
            'de la tarjeta, la cantidad de datos que proceso cada uno de sus veinte '
            'multiprocesadores, medida por los propios hilos durante la ejecucion del '
            'kernel.')

    def problema(self):
        self.titulo('2. Planteamiento del problema')
        self.parrafo(
            'El ADN se representa computacionalmente como una cadena de caracteres '
            'construida sobre un alfabeto de cuatro simbolos que corresponden a las '
            'bases nitrogenadas adenina, citosina, guanina y timina. El conteo de '
            'frecuencias sobre esa cadena es una operacion algoritmicamente trivial, '
            'pero deja de serlo cuando el volumen alcanza varios gigabytes, porque el '
            'costo deja de estar en el calculo y pasa a estar en el acceso a los datos.')
        self.parrafo(
            'Los proyectos anteriores abordaron el problema plataforma por plataforma: '
            'primero repartiendo el trabajo entre los nucleos de la CPU y despues '
            'trasladandolo a la GPU para comparar ambos resultados. El presente '
            'proyecto plantea una pregunta distinta y mas exigente: si ambas unidades '
            'de computo pueden trabajar al mismo tiempo sobre el mismo archivo y si '
            'esa colaboracion produce un beneficio.')
        self.parrafo(
            'La dificultad no esta en lanzar las dos ejecuciones a la vez, sino en '
            'repartir el trabajo entre unidades de velocidad muy distinta y no conocida '
            'de antemano. Un reparto fijo por porcentajes exige saber cuanto mas rapida '
            'es una que la otra, dato que depende del tamano del archivo, del estado de '
            'la cache del sistema y de la carga del equipo. Si el reparto se calibra '
            'mal, la unidad mas rapida termina su parte y queda ociosa esperando a la '
            'otra, y el conjunto rinde peor que la rapida trabajando sola.')
        self.parrafo(
            'El archivo utilizado sigue el formato FASTA, en el que las lineas que '
            'comienzan con el simbolo mayor que corresponden a cabeceras descriptivas '
            'y deben excluirse del conteo. Las lineas de secuencia contienen ademas la '
            'letra N para bases desconocidas y emplean minusculas como marcado de '
            'regiones repetitivas, que deben contarse igual que las mayusculas.')

    def objetivos(self):
        self.titulo('3. Objetivos')
        for texto in (
            'Desarrollar una aplicacion que procese un archivo FASTA de varios '
            'gigabytes empleando la CPU multinucleo y la GPU de forma simultanea '
            'sobre el mismo conjunto de datos.',
            'Disenar un mecanismo de reparto de carga que no dependa de conocer de '
            'antemano la velocidad relativa de cada plataforma.',
            'Demostrar de forma medible que ambas plataformas estuvieron activas al '
            'mismo tiempo y cuantificar esa superposicion.',
            'Instrumentar el uso de recursos a nivel de nucleo: cada procesador '
            'logico de la CPU y cada multiprocesador de la GPU.',
            'Medir la temperatura y el consumo de memoria de ambas plataformas '
            'durante el procesamiento.',
            'Analizar si la ejecucion conjunta resulta ventajosa frente a cada '
            'plataforma por separado e identificar el factor que lo determina.',
        ):
            self.vineta(texto)

    def conceptos(self):
        self.titulo('4. Marco conceptual')

        self.titulo('4.1 Componentes de un computador', 12, 10)
        self.parrafo(
            'Un computador se organiza alrededor de cuatro bloques funcionales. La '
            'unidad central de proceso ejecuta las instrucciones; la memoria principal '
            'almacena de forma volatil los datos y el codigo en uso; el almacenamiento '
            'secundario conserva la informacion de forma permanente; y los buses '
            'interconectan todo lo anterior. A estos se suman hoy los coprocesadores, '
            'entre ellos la unidad de procesamiento grafico, que dispone de su propia '
            'memoria y se comunica con el resto del sistema a traves del bus PCIe.')
        self.parrafo(
            'Esa separacion de memorias es determinante para este proyecto. Los datos '
            'que la GPU ha de procesar no residen en su memoria: deben viajar desde la '
            'memoria principal hasta la memoria de la tarjeta antes de poder ser '
            'tratados, y ese trayecto tiene un costo que la CPU no paga.')

        self.titulo('4.2 Proposito de una CPU frente al de una GPU', 12, 10)
        self.parrafo(
            'La CPU esta disenada para minimizar la latencia de tareas heterogeneas y '
            'con dependencias entre si. Dedica una fraccion considerable de su '
            'superficie a memorias cache, prediccion de saltos y ejecucion fuera de '
            'orden, con el proposito de que un hilo individual avance lo mas rapido '
            'posible. Dispone de pocos nucleos, pero cada uno es complejo y autonomo.')
        self.parrafo(
            'La GPU persigue el objetivo contrario: maximizar el rendimiento agregado '
            'sobre tareas homogeneas y sin dependencias. Dedica su superficie a '
            'unidades aritmeticas y organiza los hilos en grupos de treinta y dos, '
            'llamados warps, que ejecutan la misma instruccion sobre datos distintos. '
            'La tarjeta empleada en este proyecto agrupa esas unidades en veinte '
            'multiprocesadores capaces de mantener hasta mil quinientos treinta y seis '
            'hilos activos cada uno.')
        self.parrafo(
            'El conteo de frecuencias se ajusta al segundo modelo: el resultado de un '
            'caracter no condiciona el de ningun otro. Por ello la GPU resulta la '
            'plataforma natural para esta carga, mientras que la CPU aporta capacidad '
            'adicional solo en la medida en que existan recursos libres para ella.')

        self.titulo('4.3 Computador personal frente a servidor empresarial', 12, 10)
        self.parrafo(
            'Un computador personal se optimiza para la interactividad de un unico '
            'usuario: prioriza el costo, el consumo y la respuesta inmediata, y acepta '
            'que sus componentes trabajen de forma intermitente. Un servidor '
            'empresarial se optimiza para el rendimiento sostenido y la disponibilidad '
            'continua: incorpora memoria con correccion de errores, fuentes y discos '
            'redundantes, multiples zocalos de procesador y sistemas de refrigeracion '
            'dimensionados para operar sin interrupcion.')
        self.parrafo(
            'La diferencia es visible en las mediciones de este informe. El equipo '
            'empleado es un portatil, y su almacenamiento se convirtio en el factor '
            'limitante mucho antes que su capacidad de calculo. Un servidor destinado a '
            'este tipo de carga resolveria ese limite mediante arreglos de discos en '
            'paralelo o mediante la residencia de los datos en memoria, no mediante '
            'mas potencia de computo.')

        self.titulo('4.4 Supercomputadores', 12, 10)
        self.parrafo(
            'Un supercomputador es un sistema construido para obtener el maximo '
            'rendimiento agregado en problemas cientificos de gran escala. Se compone '
            'de miles de nodos de computo interconectados por una red de latencia muy '
            'baja, en su mayoria equipados con aceleradores graficos, y se acompana de '
            'sistemas de archivos paralelos, refrigeracion liquida e instalaciones '
            'electricas dedicadas. Su rendimiento se expresa en operaciones de punto '
            'flotante por segundo y los sistemas actuales de primer nivel superan la '
            'escala de exaflops.')
        self.parrafo(
            'Su proposito es abordar problemas que no admiten solucion en un plazo '
            'razonable sobre equipos convencionales: simulacion climatica, dinamica '
            'molecular, diseno de farmacos, astrofisica, ensayos nucleares simulados y, '
            'de forma creciente, entrenamiento de modelos de inteligencia artificial. '
            'La genomica figura entre sus aplicaciones habituales, y el problema '
            'tratado en este informe es una version reducida de esa clase de cargas: '
            'el mismo principio de repartir un volumen grande de datos entre multiples '
            'unidades de computo, aplicado aqui a dos unidades dentro de una sola '
            'maquina en lugar de a miles de nodos.')

    def ambiente(self):
        self.titulo('5. Ambiente de pruebas')
        self.parrafo(
            'Todas las mediciones se realizaron sobre el mismo equipo y bajo las '
            'mismas condiciones. Las especificaciones son las siguientes:')
        for texto in (
            'Procesador: Intel Core i5-13420H, 13.a generacion, arquitectura hibrida.',
            'Nucleos fisicos: 8. Procesadores logicos: 12.',
            'Tarjeta grafica: NVIDIA GeForce RTX 4050 Laptop GPU, arquitectura Ada '
            'Lovelace, compute capability 8.9, 20 multiprocesadores de flujo, 6140 MB '
            'de memoria VRAM, tamano de warp de 32 hilos, 1536 hilos y 100 KB de '
            'memoria compartida por multiprocesador.',
            'Memoria RAM: 32 GB. Almacenamiento: unidad de estado solido NVMe.',
            'Sistema operativo: Windows 11.',
            'Lenguaje: Python 3.13, con numba 0.67 y numba-cuda 0.30 para la GPU, '
            'multiprocessing para la CPU, llvmlite para la inyeccion de ensamblador '
            'PTX, y psutil, pynvml, matplotlib y tkinter para el monitoreo y la '
            'interfaz.',
            'Archivo de entrada: GCF_000001405.40_GRCh38.p14_genomic.fna, de 3,11 GB, '
            'y recortes de 50, 200 y 500 MB derivados de el.',
        ):
            self.vineta(texto)
        self.imagen(os.path.join(SALIDA, 'ui_ejecucion.png'),
                    'Aplicacion desarrollada, con el hardware detectado y el '
                    'reparto de la ultima ejecucion en el panel izquierdo.')

    def diseno(self):
        self.titulo('6. Diseno de la solucion')

        self.titulo('6.1 Un mismo algoritmo sobre las dos plataformas', 12, 10)
        self.parrafo(
            'Ambos motores calculan un histograma completo de 256 posiciones, una por '
            'cada valor posible de byte, y derivan de el todos los conteos. La decision '
            'es deliberada. Si la CPU empleara un atajo distinto al de la GPU, los '
            'tiempos dejarian de ser comparables, pero en este proyecto la consecuencia '
            'seria mas grave: el mecanismo de reparto decide a quien entregar el '
            'siguiente trozo segun la velocidad que cada plataforma demuestra, de modo '
            'que estaria repartiendo trabajos que no cuestan lo mismo y el equilibrio '
            'resultante no significaria nada.')

        self.titulo('6.2 El reparto por raciones adaptativas', 12, 10)
        self.parrafo(
            'El archivo se divide en trozos de tamano fijo numerados de forma '
            'correlativa, y un unico contador compartido indica cual es el siguiente '
            'trozo sin asignar. Cada trabajador, sea un proceso de CPU o el hilo que '
            'conduce la GPU, repite el mismo ciclo: reclama trabajo, lo procesa y '
            'vuelve a por mas hasta que se agota.')
        self.parrafo(
            'Durante el desarrollo se comprobo que pedir un trozo por vez no basta. '
            'Con doce procesos de CPU y un hilo de GPU, los trece reclaman '
            'simultaneamente en el instante inicial y se reparten trece trozos antes de '
            'que ninguno haya demostrado su velocidad; si el archivo produce dieciseis '
            'trozos, el reparto queda decidido por el orden de llegada.')
        self.parrafo(
            'La solucion adoptada consiste en que cada trabajador reclame una racion '
            'cuyo tamano ajusta segun su propia velocidad recien medida, tomando como '
            'objetivo la cantidad de trabajo que espera despachar en un cuarto de '
            'segundo. Todos comienzan pidiendo un solo trozo, en igualdad de '
            'condiciones. La GPU se estabiliza reclamando raciones amplias y los '
            'procesos de CPU se mantienen pidiendo de uno en uno. Ninguna proporcion '
            'esta escrita en el programa: emerge de la medicion, se readapta si las '
            'condiciones cambian y tiene el efecto adicional de entregar a la GPU '
            'tramos largos y contiguos, que es la forma en que mejor aprovecha la '
            'transferencia por el bus.')

        self.titulo('6.3 El descarte de cabeceras dentro de la tarjeta', 12, 10)
        self.parrafo(
            'Las cabeceras del formato FASTA no deben contarse. Eliminarlas en el '
            'equipo antes de enviar el lote a la tarjeta seria mas sencillo, pero '
            'significaria que la CPU realiza parte del trabajo de la GPU, lo que '
            'distorsionaria tanto la comparacion como el reparto de carga. Por ello el '
            'descarte ocurre dentro del dispositivo mediante tres kernels encadenados: '
            'el primero localiza el inicio de cada cabecera, el segundo la sobreescribe '
            'con saltos de linea, que el histograma ya ignora, y el tercero calcula el '
            'histograma acumulando primero en la memoria compartida de cada '
            'multiprocesador.')

        self.titulo('6.4 El arranque queda fuera de la medicion', 12, 10)
        self.parrafo(
            'La compilacion de los kernels requiere entre medio segundo y un segundo la '
            'primera vez, y la creacion de doce procesos en Windows consume un tiempo '
            'comparable. En este proyecto excluir ese costo del cronometro no es solo '
            'una cuestion de limpieza metodologica: si el arranque quedara dentro de la '
            'medicion, la plataforma que estuviera preparada antes reclamaria los '
            'primeros trozos por simple ventaja de salida, y el reparto observado no '
            'informaria sobre el rendimiento de ninguna de las dos. Por ese motivo el '
            'programa crea los procesos, los activa con tareas vacias y reserva la '
            'memoria de la tarjeta antes de iniciar el cronometro, y solo entonces da '
            'la salida a ambas plataformas de forma simultanea.')

    def implementacion(self):
        self.titulo('7. Implementacion')
        self.parrafo(
            'La solucion se estructuro en modulos independientes. El motor de CPU y el '
            'motor de GPU exponen la misma interfaz y devuelven el mismo diccionario de '
            'resultados; el motor hibrido los coordina; un modulo de monitoreo muestrea '
            'el uso de recursos en hilos de fondo; un banco de pruebas ejecuta los '
            'barridos; y la interfaz grafica reune todo lo anterior.')

        self.titulo('7.1 Codigo, libreria e instruccion de cada nivel', 12, 10)
        self.parrafo(
            'La siguiente tabla identifica, para cada uno de los tres niveles de '
            'ejecucion, el archivo que contiene el codigo correspondiente, la libreria '
            'que habilita ese nivel y la instruccion concreta que pone a trabajar al '
            'hardware.')
        self.tabla(
            ['Criterio', 'CPU', 'GPU', 'Nucleos de GPU'],
            [[nombre, cpu, gpu, nucleos] for nombre, cpu, gpu, nucleos in NIVELES])

        self.titulo('7.2 La medicion dentro de cada multiprocesador', 12, 10)
        self.parrafo(
            'Conocer el uso global de la tarjeta no permite saber como se distribuyo el '
            'trabajo entre sus multiprocesadores. Para obtener ese dato, el kernel del '
            'histograma consulta el registro especial que identifica el multiprocesador '
            'en el que se ejecuta cada hilo. Ese registro no es accesible desde Python, '
            'por lo que se recurre a un intrinsic de Numba que inyecta una instruccion '
            'de ensamblador PTX dentro del kernel compilado.')
        self.parrafo(
            'Cada hilo lleva la cuenta de los bytes que procesa en una variable local y '
            'la suma al contador de su multiprocesador una unica vez, al terminar su '
            'bucle, y no en cada iteracion. De este modo la medicion no consume tiempo '
            'del trabajo que esta midiendo, y el reparto entre los veinte '
            'multiprocesadores queda registrado por los propios hilos en lugar de '
            'estimado.')
        self.imagen(os.path.join(SALIDA, 'ui_evidencias.png'),
                    'Pestana de evidencias de la aplicacion, que reune en una sola '
                    'vista el codigo, la libreria, la instruccion y las mediciones '
                    'de cada nivel.')

        self.titulo('7.3 Concurrencia y comportamiento en Windows', 12, 10)
        self.parrafo(
            'El procesamiento se ejecuta siempre en un hilo separado del de la '
            'interfaz. Dado que Tkinter no admite el acceso a sus componentes desde '
            'varios hilos, el hilo trabajador nunca modifica un componente de forma '
            'directa: deposita mensajes en una cola que la ventana vacia '
            'periodicamente. La GPU, por su parte, se conduce desde un hilo y no desde '
            'un proceso, porque el contexto de CUDA pertenece al proceso que lo creo y '
            'no puede compartirse con un hijo.')
        self.parrafo(
            'Un detalle del comportamiento de Windows condiciono la arquitectura de '
            'archivos. En este sistema multiprocessing emplea el metodo spawn, de modo '
            'que cada proceso hijo vuelve a importar el modulo principal. Durante el '
            'desarrollo del proyecto anterior se midio que la simple presencia de la '
            'importacion del motor de GPU en el modulo principal elevaba el tiempo de '
            'una ejecucion paralela sobre un archivo de 2 MB de 1,34 a 5,51 segundos, '
            'porque los procesos hijos cargaban la infraestructura de CUDA sin llegar a '
            'utilizarla. En el presente proyecto el perjuicio seria mayor, ya que ese '
            'retraso recaeria sobre la mitad de CPU del motor justamente mientras '
            'compite por el reparto. Por ello el archivo de arranque contiene '
            'unicamente el lanzador y todos los componentes pesados se importan de '
            'forma diferida.')

    def resultados(self):
        self.titulo('8. Resultados')

        self.titulo('8.1 Conteo de nucleotidos', 12, 10)
        self.parrafo(
            'El procesamiento del archivo completo arrojo los siguientes resultados, '
            'identicos en todas las modalidades de ejecucion:')
        total = GENOMA['bases']
        self.tabla(
            ['Base', 'Ocurrencias', 'Porcentaje', 'Par complementario'],
            [['Adenina (A)', miles(GENOMA['A']),
              '%.2f %%' % (100.0 * GENOMA['A'] / total), 'Timina'],
             ['Citosina (C)', miles(GENOMA['C']),
              '%.2f %%' % (100.0 * GENOMA['C'] / total), 'Guanina'],
             ['Guanina (G)', miles(GENOMA['G']),
              '%.2f %%' % (100.0 * GENOMA['G'] / total), 'Citosina'],
             ['Timina (T)', miles(GENOMA['T']),
              '%.2f %%' % (100.0 * GENOMA['T'] / total), 'Adenina'],
             ['Total de bases', miles(total), '100,00 %', '-'],
             ['N (desconocidas)', miles(GENOMA['N']), '-', 'No aplica'],
             ['Codigos IUPAC', miles(GENOMA['iupac']), '-', 'No aplica']])
        self.parrafo(
            'Los valores obtenidos admiten una validacion independiente del programa. '
            'Dado que el ADN es una molecula de doble hebra en la que la adenina se '
            'empareja siempre con la timina y la citosina con la guanina, en un genoma '
            'completo ambas cantidades deben resultar practicamente iguales. La '
            'diferencia observada entre adenina y timina es del %.2f por ciento y la '
            'existente entre citosina y guanina del %.2f por ciento, lo que confirma la '
            'correctitud del conteo por una via biologica ajena al propio programa.'
            % (100.0 * abs(GENOMA['A'] - GENOMA['T']) / GENOMA['T'],
               100.0 * abs(GENOMA['C'] - GENOMA['G']) / GENOMA['G']))
        self.parrafo(
            'La verificacion de correctitud se realizo sobre el histograma completo de '
            '256 posiciones y no unicamente sobre las cuatro bases, y se repitio con '
            'distintas combinaciones de numero de procesos y de tamano de trozo, de '
            'manera que las uniones entre tramos consecutivos recayeran en posiciones '
            'diferentes del archivo en cada prueba. Todas las combinaciones produjeron '
            'histogramas identicos.')

        self.titulo('8.2 Trabajo simultaneo de ambas plataformas', 12, 10)
        self.parrafo(
            'La aplicacion registra el instante inicial y final de cada tramo junto '
            'con la identidad del trabajador que lo atendio, lo que permite construir '
            'la linea de tiempo de la ejecucion. Esa representacion constituye la '
            'evidencia directa de que ambas plataformas operaron de forma simultanea: '
            'si el programa se limitara a alternarlas, las bandas de CPU y de GPU no '
            'compartirian ninguna franja vertical.')
        self.imagen(os.path.join(SALIDA, 'ejecucion.png'),
                    'Tiempo por configuracion, reparto del trabajo, linea de tiempo '
                    'de la ejecucion conjunta y escalabilidad.')
        self.parrafo(
            'Sobre el genoma completo, la configuracion con cuatro procesos de CPU '
            'mantuvo ambas plataformas activas de forma simultanea durante %.2f de los '
            '%.3f segundos que duro la ejecucion, es decir durante el %.0f por ciento '
            'del tiempo total. El reparto resultante entrego el %.0f por ciento del '
            'archivo a la CPU y el resto a la GPU, proporcion que el programa no fija '
            'en ningun punto sino que deriva de la velocidad medida a cada una.'
            % (MODOS_GENOMA[2][4], MODOS_GENOMA[2][1],
               100.0 * MODOS_GENOMA[2][4] / MODOS_GENOMA[2][1],
               MODOS_GENOMA[2][3]))

        self.titulo('8.3 Desempeno comparativo', 12, 10)
        self.parrafo(
            'La tabla siguiente resume las mediciones sobre el genoma completo de '
            '3,11 GB.')
        self.tabla(
            ['Configuracion', 'Tiempo (s)', 'MB/s', 'Reparto CPU', 'Trabajo a la vez'],
            [[etiqueta, '%.3f' % tiempo, '%.1f' % velocidad,
              ('%.0f %%' % cpu) if cpu is not None else '-',
              ('%.2f s' % solape) if solape is not None else '-']
             for etiqueta, tiempo, velocidad, cpu, solape in MODOS_GENOMA])
        self.parrafo(
            'Sobre el recorte de %.0f MB, que permite repetir muchas configuraciones '
            'en un tiempo razonable, los resultados fueron los siguientes:' % MB_500)
        self.tabla(
            ['Configuracion', 'Tiempo (s)', 'MB/s', 'Speedup', 'Trabajo a la vez'],
            [[fila['etiqueta'], '%.3f' % t(fila), '%.1f' % mbs(fila),
              '%.2fx' % sp(fila),
              ('%.2f s' % sol(fila)) if fila['plataforma'] == 'HIBRIDO' else '-']
             for fila in MODOS_500])

        self.titulo('8.4 Cuantos procesos de CPU conviene sumar a la GPU', 12, 10)
        self.parrafo(
            'El numero de procesos de CPU que acompanan a la tarjeta resulto ser el '
            'parametro mas influyente del sistema, y su efecto es contrario al que '
            'cabria esperar.')
        self.tabla(
            ['Configuracion', 'Tiempo (s)', 'MB/s', 'Reparto CPU', 'Trabajo a la vez'],
            [[fila['etiqueta'], '%.3f' % t(fila), '%.1f' % mbs(fila),
              '%.1f %%' % rcpu(fila), '%.2f s' % sol(fila)]
             for fila in PROCESOS_500])
        rapida = mejor(PROCESOS_500)
        lenta = peor(PROCESOS_500)
        self.parrafo(
            'La configuracion mas rapida fue %s, con %.3f segundos, y la mas lenta %s, '
            'con %.3f segundos, es decir %.1f veces mas tiempo para el mismo trabajo. '
            'A medida que aumenta el numero de procesos, la CPU se lleva una fraccion '
            'mayor del archivo, que pasa del %.1f al %.1f por ciento, y el tiempo total '
            'empeora. El mecanismo de reparto funciona correctamente en todos los '
            'casos, en el sentido de que ninguna plataforma permanece ociosa; lo que '
            'ocurre es que el trabajo entregado a la CPU resulta mas costoso que el '
            'mismo trabajo entregado a la tarjeta. El analisis de la seccion siguiente '
            'examina la causa.'
            % (rapida['etiqueta'], t(rapida), lenta['etiqueta'], t(lenta),
               t(lenta) / t(rapida), rcpu(rapida), rcpu(lenta)))

        self.titulo('8.5 Grano del reparto', 12, 10)
        self.parrafo(
            'El tamano de los trozos en que se divide el archivo constituye la perilla '
            'de ajuste del balanceo de carga. Con trozos grandes hay pocas piezas que '
            'repartir y el resultado lo decide el orden de llegada; con trozos '
            'demasiado pequenos se paga muchas veces el costo fijo de posicionarse en '
            'el archivo.')
        self.tabla(
            ['Trozo', 'Tiempo (s)', 'MB/s', 'Reparto CPU', 'Trabajo a la vez'],
            [[fila['etiqueta'], '%.3f' % t(fila), '%.1f' % mbs(fila),
              '%.1f %%' % rcpu(fila), '%.2f s' % sol(fila)]
             for fila in TROZOS_500])
        grano = mejor(TROZOS_500)
        if grano:
            self.parrafo(
                'El mejor resultado se obtuvo con %s, con %.3f segundos. La diferencia '
                'entre el mejor y el peor grano es de %.1f por ciento, sensiblemente '
                'menor que la que provoca el numero de procesos, lo que confirma que '
                'el parametro determinante no es como se trocea el archivo sino '
                'cuantos consumidores compiten por leerlo.'
                % (grano['etiqueta'].lower(), t(grano),
                   100.0 * (t(peor(TROZOS_500)) - t(grano)) / t(grano)))

        self.titulo('8.6 Uso por nucleo de ambas plataformas', 12, 10)
        self.parrafo(
            'Durante el analisis completo, los doce procesadores logicos registraron un '
            'uso medio del %.1f por ciento, con valores comprendidos entre el %.1f y el '
            '%.1f por ciento. La dispersion refleja la arquitectura hibrida del '
            'procesador, en la que conviven nucleos de rendimiento y nucleos de '
            'eficiencia con capacidades distintas.'
            % (RECURSOS['cpu_medio'], RECURSOS['cpu_min'], RECURSOS['cpu_max']))
        self.parrafo(
            'Dentro de la tarjeta, los veinte multiprocesadores recibieron entre el '
            '%.2f y el %.2f por ciento de los bytes procesados, frente al 5,00 por '
            'ciento que corresponderia a un reparto perfectamente uniforme. La '
            'desviacion, inferior a tres decimas de punto porcentual, procede de la '
            'forma en que el planificador de la tarjeta distribuye los bloques entre '
            'los multiprocesadores y resulta despreciable a efectos practicos.'
            % (min(CENSO_SM), max(CENSO_SM)))
        self.imagen(os.path.join(SALIDA, 'nucleos.png'),
                    'Uso medio de cada procesador logico de la CPU y reparto de los '
                    'datos entre los veinte multiprocesadores de la tarjeta.')
        self.imagen(os.path.join(SALIDA, 'ui_nucleos.png'),
                    'La misma informacion dentro de la aplicacion, con las lecturas '
                    'de recursos en vivo en la barra inferior.')

        self.titulo('8.7 Temperatura y uso de memoria', 12, 10)
        self.parrafo(
            'El monitoreo registro durante el analisis una temperatura media del '
            'procesador de %.1f grados centigrados, con un maximo de %.1f, y una '
            'temperatura media de la tarjeta de %.1f grados con un maximo de %.0f. La '
            'diferencia es coherente con el reparto observado y con el hecho de que la '
            'tarjeta despacha su parte del trabajo en una fraccion del tiempo.'
            % (RECURSOS['cpu_temp_media'], RECURSOS['cpu_temp_max'],
               RECURSOS['gpu_temp_media'], RECURSOS['gpu_temp_max']))
        self.parrafo(
            'En cuanto a la memoria, el sistema alcanzo un %.1f por ciento de '
            'ocupacion y el conjunto del programa y sus procesos hijos llego a %s MB '
            'de memoria residente. La tarjeta ocupo %s MB de sus %s MB de VRAM. Se '
            'informa la memoria del programa y no la del sistema completo porque esta '
            'ultima esta dominada por el resto de aplicaciones abiertas y ocultaria por '
            'completo el costo que corresponde a cada estrategia.'
            % (RECURSOS['ram_pct'], miles(RECURSOS['proc_mb']),
               miles(RECURSOS['vram_mb']), miles(RECURSOS['vram_total'])))
        self.parrafo(
            'A nivel de multiprocesador, el kernel del histograma consume %d registros '
            'por hilo y %s bytes de memoria compartida por bloque. Con ese consumo la '
            'tarjeta admite %d de los %d warps que un multiprocesador puede mantener '
            'activos, es decir una ocupacion del %.0f por ciento, que es el valor '
            'maximo alcanzable.'
            % (RECURSOS['registros'], miles(RECURSOS['compartida']),
               RECURSOS['warps'], RECURSOS['warps_max'], RECURSOS['ocupacion']))

        self.titulo('8.8 Escalabilidad', 12, 10)
        self.parrafo(
            'El ultimo barrido midio el comportamiento de las tres modalidades frente a '
            'volumenes crecientes de datos.')
        self.tabla(
            ['Tamano', 'CPU (s)', 'GPU sola (s)', 'Conjunta (s)',
             'GPU frente a CPU'],
            [[etiqueta,
              '%.3f' % tiempos.get('CPU', 0.0),
              '%.3f' % tiempos.get('GPU', 0.0),
              '%.3f' % tiempos.get('HIBRIDO', 0.0),
              '%.2fx' % (tiempos.get('CPU', 0.0) / tiempos['GPU'])
              if tiempos.get('GPU') else '-']
             for etiqueta, tiempos in ESCALABILIDAD])
        self.parrafo(
            'La ventaja de la tarjeta se mantiene en todo el rango medido y crece con '
            'el volumen, lo que confirma que su costo fijo de preparacion se amortiza '
            'ya en los tamanos mas pequenos de esta serie. La modalidad conjunta se '
            'situa de forma sistematica entre ambas, sin alcanzar en ningun tamano a '
            'la tarjeta trabajando sola.')
        self.imagen(os.path.join(SALIDA, 'ui_tabla.png'),
                    'Tabla comparativa de la aplicacion con todas las '
                    'configuraciones medidas y la verificacion de conteos.')

    def analisis(self):
        self.titulo('9. Analisis critico')

        self.titulo('9.1 La colaboracion funciona pero no compensa', 12, 10)
        self.parrafo(
            'El objetivo central del proyecto se cumplio: ambas plataformas procesan el '
            'mismo archivo de forma simultanea, la superposicion se mide en segundos y '
            'el reparto se equilibra por si solo sin que ninguna proporcion este fijada '
            'en el programa. Sin embargo, el conjunto resulta mas lento que la tarjeta '
            'trabajando sola, y cuantos mas procesos de CPU se incorporan, peor es el '
            'resultado. Sobre el genoma completo la GPU sola tarda %.3f segundos '
            'mientras que la ejecucion conjunta con doce procesos tarda %.3f.'
            % (MODOS_GENOMA[1][1], MODOS_GENOMA[4][1]))
        self.parrafo(
            'Conviene subrayar que esto no constituye un defecto del mecanismo de '
            'reparto. El reparto hace exactamente lo que debe: ninguna plataforma '
            'permanece ociosa y cada una recibe una cantidad de trabajo proporcional a '
            'la velocidad que demuestra. El problema es de otra naturaleza y se examina '
            'a continuacion.')

        self.titulo('9.2 El cuello de botella no esta en el calculo', 12, 10)
        self.parrafo(
            'Contar bytes es una operacion de coste computacional minimo: una lectura '
            'de memoria y un incremento por cada caracter. En consecuencia, el factor '
            'que limita el rendimiento no es la capacidad de calculo de ninguna de las '
            'dos plataformas sino la velocidad a la que los datos llegan desde el '
            'almacenamiento. La GPU por si sola alcanza los %.1f MB/s sobre el genoma '
            'completo, cifra que corresponde al limite practico de lectura del sistema.'
            % (MODOS_GENOMA[1][2]))
        self.parrafo(
            'Los procesos de CPU no aportan capacidad de lectura adicional; compiten '
            'por la que ya existe. Peor aun, restan turnos de acceso al unico hilo '
            'lector que alimenta a la tarjeta, que es precisamente el que obtiene mayor '
            'rendimiento de cada byte recibido. De ahi la progresion observada: con dos '
            'procesos la perdida es moderada, mientras que con doce el conjunto rinde a '
            'menos de un tercio de lo que rendiria la tarjeta sola.')
        self.parrafo(
            'La conclusion general que se desprende es que la ejecucion conjunta de dos '
            'unidades de computo solo resulta ventajosa cuando el factor limitante es '
            'el computo. Cuando el limite reside en el acceso a los datos, ningun '
            'reparto, por perfecto que sea, puede superar el ancho de banda del '
            'componente mas lento de la cadena, y anadir consumidores de ese recurso '
            'escaso lo unico que consigue es repartirlo peor.')

        self.titulo('9.3 El reparto dentro de la tarjeta', 12, 10)
        self.parrafo(
            'Frente al desequilibrio observado entre las dos plataformas, el reparto '
            'interno de la tarjeta resulto notablemente uniforme: los veinte '
            'multiprocesadores procesaron entre el %.2f y el %.2f por ciento de los '
            'datos. El contraste es ilustrativo, porque ambos repartos obedecen a '
            'mecanismos distintos. El reparto entre multiprocesadores lo realiza el '
            'planificador de la tarjeta sobre unidades identicas y con acceso a la '
            'misma memoria, situacion en la que un reparto uniforme es el resultado '
            'esperable. El reparto entre CPU y GPU, en cambio, se produce entre '
            'unidades de naturaleza distinta que compiten por un recurso externo y '
            'limitado.'
            % (min(CENSO_SM), max(CENSO_SM)))

        self.titulo('9.4 Un error de borde en las uniones entre tramos', 12, 10)
        self.parrafo(
            'Durante la verificacion se detecto que el conteo resultaba incorrecto en '
            'aproximadamente una de cada cinco ejecuciones, sin patron aparente. El '
            'origen estaba en el tratamiento del final de cada tramo por parte del '
            'motor de GPU: al cerrar el ultimo lote llamaba a la funcion de lectura de '
            'linea sin comprobar previamente si la linea habia quedado efectivamente '
            'cortada. Cuando un tramo terminaba justo en un salto de linea, esa llamada '
            'devolvia la linea siguiente completa, que pertenecia al tramo contiguo, y '
            'esos ochenta caracteres acababan contados dos veces.')
        self.parrafo(
            'El fallo permanecia oculto por dos razones. Las divisiones se producen '
            'cada dieciseis megabytes y las lineas miden ochenta y un bytes, de modo '
            'que la coincidencia es infrecuente; y ademas solo se manifestaba cuando el '
            'tramo que terminaba en un salto correspondia a la GPU y no a la CPU, que '
            'si realizaba la comprobacion. Como el reparto varia en cada ejecucion, el '
            'error aparecia y desaparecia sin causa aparente.')
        self.parrafo(
            'El episodio deja una leccion metodologica aplicable a cualquier sistema '
            'con reparto dinamico: una unica ejecucion correcta no demuestra nada, '
            'porque las uniones entre tramos recaen en posiciones distintas en cada '
            'una. Por ese motivo el programa de verificacion ejecuta varias '
            'combinaciones de numero de procesos y de tamano de trozo, con el proposito '
            'expreso de desplazar esas uniones a lo largo del archivo.')

    def aplicaciones(self):
        self.titulo('10. Aplicaciones reales')
        self.parrafo(
            'El escenario planteado constituye una version simplificada de problemas '
            'habituales en bioinformatica. El alineamiento de secuencias contra un '
            'genoma de referencia, la deteccion de variantes geneticas, el calculo del '
            'contenido de guanina y citosina por regiones o el control de calidad de '
            'los datos producidos por un secuenciador operan sobre volumenes '
            'comparables o superiores.')
        self.parrafo(
            'La ejecucion conjunta sobre plataformas heterogeneas es hoy la norma en '
            'los sistemas de alto rendimiento. Los entornos de computo cientifico '
            'reparten el trabajo entre procesadores y aceleradores, y los marcos de '
            'inteligencia artificial distribuyen las operaciones entre ambos segun su '
            'naturaleza. El mecanismo de reparto dinamico implementado en este proyecto '
            'es una version reducida del que emplean esos sistemas.')
        self.parrafo(
            'El resultado obtenido tiene ademas valor practico mas alla del caso '
            'concreto. Antes de invertir esfuerzo en repartir una carga entre varias '
            'unidades de computo conviene determinar donde reside el factor limitante. '
            'Si se encuentra en el acceso a los datos, la solucion no pasa por sumar '
            'unidades de calculo sino por mejorar el camino de los datos: '
            'almacenamiento mas rapido, lectura en paralelo desde varios dispositivos o '
            'residencia de los datos en memoria.')

    def conclusiones(self):
        self.titulo('11. Conclusiones')
        for texto in (
            'Se implemento un motor que procesa un mismo archivo con la CPU '
            'multinucleo y la GPU trabajando de forma simultanea. Sobre el genoma '
            'completo ambas plataformas permanecieron activas al mismo tiempo durante '
            '%.2f de los %.3f segundos de la ejecucion.'
            % (MODOS_GENOMA[2][4], MODOS_GENOMA[2][1]),
            'El reparto de carga se resuelve mediante raciones adaptativas: cada '
            'trabajador ajusta la cantidad de trabajo que reclama segun su propia '
            'velocidad medida. Ninguna proporcion entre plataformas esta fijada en el '
            'programa, y el reparto obtenido es una consecuencia medida y no una '
            'suposicion.',
            'La ejecucion conjunta resulto mas lenta que la GPU trabajando sola: %.3f '
            'segundos frente a %.3f sobre el genoma completo en la mejor configuracion '
            'conjunta. El factor determinante es que el problema esta limitado por el '
            'acceso a los datos y no por la capacidad de calculo.'
            % (MODOS_GENOMA[2][1], MODOS_GENOMA[1][1]),
            'El numero de procesos de CPU que acompanan a la tarjeta resulto el '
            'parametro mas influyente: entre la mejor y la peor configuracion del '
            'barrido hay un factor de %.1f para el mismo trabajo, porque todos ellos '
            'compiten por el mismo ancho de banda de lectura.'
            % (t(peor(PROCESOS_500)) / t(mejor(PROCESOS_500))),
            'La instrumentacion a nivel de nucleo mostro un reparto interno de la '
            'tarjeta muy uniforme, entre el %.2f y el %.2f por ciento de los datos por '
            'multiprocesador, y una ocupacion del %.0f por ciento de los warps '
            'disponibles.'
            % (min(CENSO_SM), max(CENSO_SM), RECURSOS['ocupacion']),
            'La correctitud se verifico sobre el histograma completo de 256 valores de '
            'byte y con multiples configuraciones de reparto, de modo que las uniones '
            'entre tramos recayeran en posiciones distintas. Todas las combinaciones '
            'produjeron conteos identicos.',
            'La proporcion obtenida entre bases complementarias, con diferencias '
            'inferiores al 0,31 por ciento entre adenina y timina y al 0,42 por ciento '
            'entre citosina y guanina, valida el resultado por una via biologica '
            'independiente del programa.',
        ):
            self.vineta(texto)

    def recomendaciones(self):
        self.titulo('12. Recomendaciones')
        for texto in (
            'Evaluar la lectura del archivo mediante varios hilos de entrada y salida '
            'con profundidad de cola elevada, dado que las unidades NVMe admiten '
            'multiples peticiones simultaneas y podrian superar el limite alcanzado '
            'por el acceso secuencial.',
            'Repetir el experimento con una carga de mayor costo computacional por '
            'byte, como la busqueda de patrones o el calculo de contenido de guanina '
            'y citosina por ventanas, para verificar la hipotesis de que la ejecucion '
            'conjunta resulta ventajosa cuando el limite reside en el computo.',
            'Incorporar al mecanismo de reparto una fase inicial de calibracion que '
            'mida ambas plataformas por separado antes de comenzar, de modo que las '
            'primeras raciones partan de una estimacion mas ajustada.',
            'Repetir cada configuracion varias veces y reportar el promedio, con el fin '
            'de reducir el efecto de la cache del sistema operativo, cuya influencia '
            'sobre las cifras absolutas resulto considerable.',
            'Considerar la implementacion del motor de CPU en un lenguaje compilado '
            'para establecer el limite superior real de esa plataforma y obtener una '
            'comparacion mas estricta frente a la tarjeta.',
        ):
            self.vineta(texto)

    def referencias(self):
        self.titulo('13. Referencias')
        for texto in (
            'NVIDIA Corporation. CUDA C++ Programming Guide. Disponible en: '
            'https://docs.nvidia.com/cuda/cuda-c-programming-guide/',
            'NVIDIA Corporation. Parallel Thread Execution ISA. Disponible en: '
            'https://docs.nvidia.com/cuda/parallel-thread-execution/',
            'Anaconda Inc. Numba Documentation: CUDA Python. Disponible en: '
            'https://numba.readthedocs.io/en/stable/cuda/index.html',
            'Python Software Foundation. Documentacion oficial del modulo '
            'multiprocessing. Disponible en: '
            'https://docs.python.org/3/library/multiprocessing.html',
            'National Center for Biotechnology Information. Genome Reference '
            'Consortium Human Build 38 (GRCh38). Disponible en: '
            'https://www.ncbi.nlm.nih.gov/assembly/GCF_000001405.40/',
            'Cornish-Bowden, A. Nomenclature for incompletely specified bases in '
            'nucleic acid sequences. Nucleic Acids Research, vol. 13, num. 9, 1985.',
            'TOP500.org. TOP500 List of the most powerful supercomputer systems. '
            'Disponible en: https://www.top500.org/',
            'NVIDIA Corporation. Especificaciones de la arquitectura Ada Lovelace y '
            'de la serie GeForce RTX 40. Disponible en: https://www.nvidia.com/',
        ):
            self.vineta(texto)

    def generar(self):
        self.portada()
        self.doc.add_page_break()
        self.resumen()
        self.problema()
        self.objetivos()
        self.conceptos()
        self.ambiente()
        self.diseno()
        self.implementacion()
        self.resultados()
        self.analisis()
        self.aplicaciones()
        self.conclusiones()
        self.recomendaciones()
        self.referencias()
        self.doc.save(DESTINO)
        print('Informe generado: %s' % os.path.abspath(DESTINO))
        print('Figuras incluidas: %d' % self.figura)


if __name__ == '__main__':
    Informe().generar()
